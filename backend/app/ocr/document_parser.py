import io
from pathlib import Path
from app.ocr.pdf_parser import extract_text_from_pdf
from app.ocr.image_ocr import extract_text_from_image

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        paragraphs.extend(' | '.join(cell.text.strip() for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n\n".join(paragraphs)
    except Exception as e:
        print(f"[DocumentParser] Error parsing docx: {e}")
        return ""

def extract_text_from_txt(file_bytes: bytes) -> str:
    for encoding in ["utf-8", "latin-1", "windows-1252"]:
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()

def extract_document_text(filename: str, file_bytes: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(file_bytes)
        if not text:
            # try OCR as fallback for scanned PDF
            text = extract_text_from_image(file_bytes)
        return text
    elif ext == ".docx":
        return extract_text_from_docx(file_bytes)
    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp"]:
        return extract_text_from_image(file_bytes)
    elif ext in [".txt", ".text"]:
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX, TXT, PNG, JPG, WEBP or BMP.")
