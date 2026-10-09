import io
import docx
import pymupdf
import pytest
from app.ocr.document_parser import extract_document_text

def test_digital_pdf_docx_tables_and_text():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), 'Students must register for examination tomorrow.')
    assert 'examination' in extract_document_text('notice.pdf', pdf.tobytes())
    pdf.close()
    document = docx.Document()
    document.add_paragraph('Workshop registration')
    document.add_table(rows=1, cols=1).cell(0,0).text = 'Deadline tomorrow'
    stream = io.BytesIO()
    document.save(stream)
    assert 'Deadline tomorrow' in extract_document_text('notice.docx', stream.getvalue())
    assert extract_document_text('notice.txt', b'Exam registration') == 'Exam registration'

def test_unsupported_document_is_rejected():
    with pytest.raises(ValueError, match='Unsupported'):
        extract_document_text('file.exe', b'not a notice')

def test_failed_ocr_never_invents_text(monkeypatch):
    import app.ocr.image_ocr as ocr
    monkeypatch.setattr(ocr, 'get_ocr_engine', lambda: False)
    with pytest.raises(ValueError, match='OCR'):
        ocr.extract_text_from_image(b'invalid image')

def test_actions_keep_complete_requirements():
    from app.ai.entity_extractor import extract_actions
    assert extract_actions('Students must carry their college ID cards and hall tickets.') == ['Carry college ID cards and hall tickets']
